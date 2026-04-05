import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ---------- Read data ----------
ppu_all = pd.read_csv('/home/accelgen/test_ppu/ppu_detail.csv')
edge_raw = pd.read_csv('/home/accelgen/edge_results.csv')

columns = list(ppu_all.columns)  # 23 columns

# ---------- PPU edge only ----------
ppu_edge = ppu_all[ppu_all['Config'] == 'edge'].copy()

# ---------- Normalize edge_results ----------
edge = edge_raw.copy()
# Arch: gemmini -> gemmini_edge, lego -> lego_edge
edge['Arch'] = edge['Arch'].apply(lambda x: x + '_edge' if not x.endswith('_edge') else x)
# Model: gemma_7b -> gemma-7b, etc.
edge['Model'] = edge['Model'].str.replace('_', '-')
# Derive normalized layer and keep original Layer as part of OpName
# attn_prefill_0 -> layer=attention, attn_decode_3 -> layer=attention
# ffn_prefill_1 -> layer=ffn
edge['_norm_layer'] = edge['Layer'].apply(
    lambda x: 'attention' if x.startswith('attn') else 'ffn'
)
# Phase is already correct in edge_results

# ---------- Define workload order ----------
archs = ['gemmini_edge', 'lego_edge']
models = ['llama3-8b', 'gemma-7b', 'qwen3-8b']
layers = ['attention', 'ffn']
phases = ['prefill', 'decode']

# ---------- Numeric columns for summing ----------
sum_cols = ['Cycles', 'Energy(uJ)', 'Computes', 'DRAM_Acc', 'SRAM_Acc']

# ---------- Build workbook ----------
wb = Workbook()

# ===== Sheet 1: Edge_Complete (grouped workloads) =====
ws1 = wb.active
ws1.title = 'Edge_Complete'

# Styles
header_font = Font(bold=True)
total_font = Font(bold=True, color='FFFFFF')
total_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
group_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')

# Write header
for ci, col_name in enumerate(columns, 1):
    cell = ws1.cell(row=1, column=ci, value=col_name)
    cell.font = header_font

row_ptr = 2  # current row pointer
workload_count = 0

for arch in archs:
    for model in models:
        for layer in layers:
            for phase in phases:
                workload_count += 1

                # Get linear ops from edge_results for this arch/model/layer/phase
                linear_ops = edge[
                    (edge['Arch'] == arch) &
                    (edge['Model'] == model) &
                    (edge['_norm_layer'] == layer) &
                    (edge['Phase'] == phase)
                ].copy()
                # Sort by original Layer name (attn_prefill_0, 1, 2...)
                linear_ops = linear_ops.sort_values('Layer')

                # Get PPU elem+datamove ops for this model/layer/phase
                ppu_ops = ppu_edge[
                    (ppu_edge['Model'] == model) &
                    (ppu_edge['Layer'] == layer) &
                    (ppu_edge['Phase'] == phase)
                ].copy()

                # --- Write workload group label row ---
                label = f"=== {arch} / {model} / {layer} / {phase} ==="
                cell = ws1.cell(row=row_ptr, column=1, value=label)
                cell.font = Font(bold=True, size=11)
                for ci in range(1, len(columns) + 1):
                    ws1.cell(row=row_ptr, column=ci).fill = group_fill
                row_ptr += 1

                # Collect all ops for total calculation
                all_ops_for_total = []

                # --- Write linear ops (from gemmini/lego) ---
                for _, op_row in linear_ops.iterrows():
                    for ci, col_name in enumerate(columns, 1):
                        if col_name == 'Layer':
                            # Use normalized layer name
                            ws1.cell(row=row_ptr, column=ci, value=layer)
                        elif col_name == 'OpName':
                            # Use original Layer value as OpName (attn_prefill_0, etc.)
                            ws1.cell(row=row_ptr, column=ci, value=op_row['Layer'])
                        elif col_name == 'OpType':
                            ws1.cell(row=row_ptr, column=ci, value='linear')
                        else:
                            val = op_row.get(col_name, '')
                            ws1.cell(row=row_ptr, column=ci, value=val if pd.notna(val) else '')
                    all_ops_for_total.append(op_row)
                    row_ptr += 1

                # --- Write PPU elem+datamove ops ---
                for _, op_row in ppu_ops.iterrows():
                    for ci, col_name in enumerate(columns, 1):
                        val = op_row.get(col_name, '')
                        ws1.cell(row=row_ptr, column=ci, value=val if pd.notna(val) else '')
                    all_ops_for_total.append(op_row)
                    row_ptr += 1

                # --- Write TOTAL row ---
                for ci, col_name in enumerate(columns, 1):
                    cell = ws1.cell(row=row_ptr, column=ci)
                    cell.font = total_font
                    cell.fill = total_fill

                    if col_name == 'Arch':
                        cell.value = arch
                    elif col_name == 'Config':
                        cell.value = 'edge'
                    elif col_name == 'Model':
                        cell.value = model
                    elif col_name == 'Layer':
                        cell.value = layer
                    elif col_name == 'Phase':
                        cell.value = phase
                    elif col_name == 'OpName':
                        cell.value = 'TOTAL'
                    elif col_name == 'OpType':
                        cell.value = ''
                    elif col_name in sum_cols:
                        total_val = 0
                        for op in all_ops_for_total:
                            v = op.get(col_name, 0)
                            if pd.notna(v):
                                try:
                                    total_val += float(v)
                                except (ValueError, TypeError):
                                    pass
                        # Keep as int for integer columns
                        if col_name in ['Cycles', 'Computes', 'DRAM_Acc', 'SRAM_Acc']:
                            cell.value = int(total_val)
                        else:
                            cell.value = round(total_val, 6)
                    else:
                        cell.value = ''
                row_ptr += 1

                # --- 2 blank rows ---
                row_ptr += 2

print(f"Sheet 'Edge_Complete': {workload_count} workload blocks, last row = {row_ptr - 1}")

# ===== Sheet 2: PPU_Edge_Detail =====
ws2 = wb.create_sheet('PPU_Edge_Detail')
for ci, col_name in enumerate(columns, 1):
    ws2.cell(row=1, column=ci, value=col_name).font = header_font

for ri, (_, row_data) in enumerate(ppu_edge.iterrows(), 2):
    for ci, col_name in enumerate(columns, 1):
        val = row_data.get(col_name, '')
        ws2.cell(row=ri, column=ci, value=val if pd.notna(val) else '')

print(f"Sheet 'PPU_Edge_Detail': {len(ppu_edge)} rows")

# ===== Sheet 3: Gemmini_Lego_Edge =====
ws3 = wb.create_sheet('Gemmini_Lego_Edge')
for ci, col_name in enumerate(columns, 1):
    ws3.cell(row=1, column=ci, value=col_name).font = header_font

for ri, (_, row_data) in enumerate(edge_raw.iterrows(), 2):
    for ci, col_name in enumerate(columns, 1):
        val = row_data.get(col_name, '')
        ws3.cell(row=ri, column=ci, value=val if pd.notna(val) else '')

print(f"Sheet 'Gemmini_Lego_Edge': {len(edge_raw)} rows")

# ===== Auto-fit column widths for all sheets =====
for ws in [ws1, ws2, ws3]:
    for ci in range(1, len(columns) + 1):
        max_len = len(str(ws.cell(row=1, column=ci).value or ''))
        for ri in range(2, min(ws.max_row + 1, 50)):  # sample first 50 rows
            val = ws.cell(row=ri, column=ci).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 40)

# Save
out_path = '/home/accelgen/test_ppu/edge_workload_summary.xlsx'
wb.save(out_path)
print(f"\nSaved to {out_path}")
