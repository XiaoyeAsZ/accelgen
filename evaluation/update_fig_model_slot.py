#!/usr/bin/env python3
"""Replace the first model slot in evaluation/fig workbooks with Llama3-70B."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


MODEL = "llama3-70b"
MODELS = ("llama3-70b", "qwen3-8b", "gemma-7b", "qwen3-moe")
LENGTHS = (128, 256, 512, 1024, 4096)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def index(rows: list[dict[str, str]], *fields: str) -> dict[tuple[str, ...], dict[str, str]]:
    return {tuple(row[field] for field in fields): row for row in rows}


def value(row: dict[str, str], field: str) -> float:
    return float(row[field])


def set_recalculate(workbook) -> None:
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"


def update_standard(
    workbook_path: Path,
    baseline: dict[tuple[str, str, str, str], dict[str, str]],
    ours: dict[tuple[str, str, str], dict[str, str]],
) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    for row_number, length in enumerate(LENGTHS, 2):
        metrics = [
            baseline[("gemmini_os", "edge" if "edge" in workbook_path.name else "server", "prefill" if "prefill" in workbook_path.name else "decode", str(length))],
            baseline[("gemmini_ws", "edge" if "edge" in workbook_path.name else "server", "prefill" if "prefill" in workbook_path.name else "decode", str(length))],
            baseline[("lego", "edge" if "edge" in workbook_path.name else "server", "prefill" if "prefill" in workbook_path.name else "decode", str(length))],
        ]
        own = ours[("prefill" if "prefill" in workbook_path.name else "decode", "edge" if "edge" in workbook_path.name else "server", str(length))]
        throughput = [value(metric, "throughput_GFLOPS") for metric in metrics] + [value(own, "throughput_GFLOPS")]
        efficiency = [value(metric, "energy_eff_GFLOPS_per_J") for metric in metrics] + [value(own, "energy_efficiency_GFLOPS_per_J")]
        for column, metric in zip(range(2, 6), throughput):
            sheet.cell(row_number, column, metric)
        for column, metric in zip(range(26, 30), efficiency):
            sheet.cell(row_number, column, metric)
    set_recalculate(workbook)
    workbook.save(workbook_path)


def update_gpu_edge(
    workbook_path: Path,
    gpu: dict[tuple[str, str, str, str], dict[str, str]],
    orin_scaled: dict[tuple[str, str, str], dict[str, str]],
) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    action = "prefill" if "prefill" in workbook_path.name else "decode"
    for row_number, length in enumerate(LENGTHS[:-1], 2):
        device = gpu[("Orin", action, "edge", str(length))]
        own = orin_scaled[(action, "orin", str(length))]
        for column, metric in zip((2, 3), (value(device, "throughput_GFLOPS"), value(own, "throughput_GFLOPS"))):
            sheet.cell(row_number, column, metric)
        for column, metric in zip((19, 20), (value(device, "energy_efficiency_GFLOPS_per_J"), value(own, "energy_efficiency_GFLOPS_per_J"))):
            sheet.cell(row_number, column, metric)
    set_recalculate(workbook)
    workbook.save(workbook_path)


def update_gpu_server(
    workbook_path: Path,
    gpu: dict[tuple[str, str, str, str], dict[str, str]],
    ascend: dict[tuple[str, str, str], dict[str, str]],
    ours: dict[tuple[str, str, str], dict[str, str]],
) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    action = "prefill" if "prefill" in workbook_path.name else "decode"
    first_data_row = 31 if action == "prefill" else 2
    efficiency_start = 25 if action == "prefill" else 23
    for row_number, length in enumerate(LENGTHS[:-1], first_data_row):
        a100 = gpu[("A100", action, "server", str(length))]
        npu = ascend[(action, "server", str(length))]
        own = ours[(action, "server", str(length))]
        throughput = [value(a100, "throughput_GFLOPS"), value(npu, "throughput_GFLOPS"), value(own, "throughput_GFLOPS")]
        efficiency = [value(a100, "energy_efficiency_GFLOPS_per_J"), value(npu, "energy_efficiency_GFLOPS_per_J"), value(own, "energy_efficiency_GFLOPS_per_J")]
        for column, metric in zip((2, 3, 4), throughput):
            sheet.cell(row_number, column, metric)
        for column, metric in zip((efficiency_start, efficiency_start + 1, efficiency_start + 2), efficiency):
            sheet.cell(row_number, column, metric)
    set_recalculate(workbook)
    workbook.save(workbook_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fig-dir", type=Path, default=Path("evaluation/fig"))
    parser.add_argument("--baseline-csv", type=Path, default=Path("evaluation/results/llama3-70b-baseline/llama3_70b_baseline_combined.csv"))
    parser.add_argument("--model-performance-csv", type=Path, default=Path("evaluation/results/model-performance/model_performance.csv"))
    parser.add_argument("--gpu-csv", type=Path, default=Path("evaluation/results/gpu-baselines/gpu_baselines.csv"))
    parser.add_argument("--ascend-csv", type=Path, default=Path("evaluation/results/ascend-910b3/ascend_910b3_combined.csv"))
    parser.add_argument("--orin-scaled-csv", type=Path, default=Path("evaluation/results/model-performance-orin/model_performance_orin_scaled.csv"))
    args = parser.parse_args()

    baseline_rows = read_csv(args.baseline_csv)
    model_rows = read_csv(args.model_performance_csv)
    gpu_rows = read_csv(args.gpu_csv)
    ascend_rows = read_csv(args.ascend_csv)
    orin_rows = read_csv(args.orin_scaled_csv)

    baseline = index(baseline_rows, "arch", "config", "action", "length")
    ours = index([row for row in model_rows if row["model"] == MODEL], "action", "config", "length")
    gpu = index([row for row in gpu_rows if row["model"] == MODEL], "hardware", "action", "config", "length")
    ascend = index([row for row in ascend_rows if row["model"] == MODEL], "action", "config", "length")
    orin_scaled = index([row for row in orin_rows if row["model"] == MODEL], "action", "config", "length")

    for workbook_path in sorted(args.fig_dir.glob("*.xlsx")):
        if workbook_path.name.endswith("_gpu_ori.xlsx"):
            if "edge" in workbook_path.name:
                update_gpu_edge(workbook_path, gpu, orin_scaled)
            else:
                update_gpu_server(workbook_path, gpu, ascend, ours)
        else:
            update_standard(workbook_path, baseline, ours)
    print(f"Updated first model slot in {len(list(args.fig_dir.glob('*.xlsx')))} workbooks")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
