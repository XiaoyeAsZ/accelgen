#!/usr/bin/env python3
"""Create appendix-ready Llama3-70B performance tables from a model log."""

from __future__ import annotations

import argparse
import csv
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


RUN_RE = re.compile(
    r"^Running: model=(?P<model>\S+) action=(?P<action>prefill|decode) "
    r"block=(?P<block>\d+) layer=(?P<layer>attention|ffn) "
    r"batch=(?P<batch>\d+) length=(?P<length>\d+) config=(?P<config>edge|server)$"
)
LATENCY_RE = re.compile(r"^Latency:\s*(?P<value>[-+0-9.eE]+)$")
ENERGY_RE = re.compile(r"^Energy:\s*(?P<value>[-+0-9.eE]+)$")
FLOPS_RE = re.compile(r"^Flops:\s*(?P<value>[-+0-9.eE]+)$")


@dataclass(frozen=True)
class LayerRecord:
    model: str
    block: int
    action: str
    layer: str
    batch: int
    length: int
    config: str
    latency_ms: float
    energy_j: float
    flops: float


@dataclass(frozen=True)
class CombinedRecord:
    model: str
    block: int
    action: str
    batch: int
    length: int
    config: str
    attention_latency_ms: float
    ffn_latency_ms: float
    attention_energy_j: float
    ffn_energy_j: float
    attention_flops: float
    ffn_flops: float

    @property
    def latency_ms(self) -> float:
        return self.attention_latency_ms + self.ffn_latency_ms

    @property
    def energy_j(self) -> float:
        return self.attention_energy_j + self.ffn_energy_j

    @property
    def flops(self) -> float:
        return self.attention_flops + self.ffn_flops

    @property
    def throughput_gflops(self) -> float:
        return self.flops / (self.latency_ms * 1e6)

    @property
    def efficiency_gflops_per_j(self) -> float:
        return self.flops / (self.energy_j * 1e9)


def parse_log(path: Path) -> list[LayerRecord]:
    records: list[LayerRecord] = []
    current: dict[str, str] | None = None
    latency_ms = energy_j = flops = None
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        run = RUN_RE.match(line)
        if run:
            if current is not None:
                raise ValueError(f"incomplete record before {path}:{line_number}")
            current = run.groupdict()
            latency_ms = energy_j = flops = None
            continue
        if current is None:
            continue
        latency_match = LATENCY_RE.match(line)
        if latency_match:
            latency_ms = float(latency_match["value"])
            continue
        energy_match = ENERGY_RE.match(line)
        if energy_match:
            energy_j = float(energy_match["value"])
            continue
        flops_match = FLOPS_RE.match(line)
        if flops_match:
            flops = float(flops_match["value"])
            continue
        if line.startswith("============================="):
            if None in (latency_ms, energy_j, flops):
                raise ValueError(f"missing metric at {path}:{line_number}")
            records.append(
                LayerRecord(
                    model=current["model"],
                    block=int(current["block"]),
                    action=current["action"],
                    layer=current["layer"],
                    batch=int(current["batch"]),
                    length=int(current["length"]),
                    config=current["config"],
                    latency_ms=latency_ms,
                    energy_j=energy_j,
                    flops=flops,
                )
            )
            current = None
    if current is not None:
        raise ValueError(f"unterminated record at end of {path}")
    return records


def combine(records: list[LayerRecord]) -> list[CombinedRecord]:
    grouped: dict[
        tuple[str, int, str, int, int, str], dict[str, LayerRecord]
    ] = {}
    for record in records:
        key = (
            record.model,
            record.block,
            record.action,
            record.batch,
            record.length,
            record.config,
        )
        layers = grouped.setdefault(key, {})
        if record.layer in layers:
            raise ValueError(f"duplicate {record.layer} record for {key}")
        layers[record.layer] = record

    rows: list[CombinedRecord] = []
    for key, layers in grouped.items():
        if set(layers) != {"attention", "ffn"}:
            raise ValueError(f"missing attention or FFN record for {key}: {set(layers)}")
        model, block, action, batch, length, config = key
        attention = layers["attention"]
        ffn = layers["ffn"]
        rows.append(
            CombinedRecord(
                model=model,
                block=block,
                action=action,
                batch=batch,
                length=length,
                config=config,
                attention_latency_ms=attention.latency_ms,
                ffn_latency_ms=ffn.latency_ms,
                attention_energy_j=attention.energy_j,
                ffn_energy_j=ffn.energy_j,
                attention_flops=attention.flops,
                ffn_flops=ffn.flops,
            )
        )
    config_order = {"edge": 0, "server": 1}
    action_order = {"prefill": 0, "decode": 1}
    return sorted(
        rows,
        key=lambda row: (
            config_order[row.config], action_order[row.action], row.length
        ),
    )


def write_csv(rows: list[CombinedRecord], path: Path) -> None:
    fields = [
        "model", "block", "action", "config", "batch", "length",
        "attention_latency_ms", "ffn_latency_ms", "total_latency_ms",
        "attention_energy_j", "ffn_energy_j", "total_energy_j",
        "attention_flops", "ffn_flops", "total_flops", "throughput_GFLOPS",
        "energy_efficiency_GFLOPS_per_J",
    ]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "model": row.model,
                    "block": row.block,
                    "action": row.action,
                    "config": row.config,
                    "batch": row.batch,
                    "length": row.length,
                    "attention_latency_ms": row.attention_latency_ms,
                    "ffn_latency_ms": row.ffn_latency_ms,
                    "total_latency_ms": row.latency_ms,
                    "attention_energy_j": row.attention_energy_j,
                    "ffn_energy_j": row.ffn_energy_j,
                    "total_energy_j": row.energy_j,
                    "attention_flops": row.attention_flops,
                    "ffn_flops": row.ffn_flops,
                    "total_flops": row.flops,
                    "throughput_GFLOPS": row.throughput_gflops,
                    "energy_efficiency_GFLOPS_per_J": row.efficiency_gflops_per_j,
                }
            )


def apply_table_border(sheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Side(style="thin", color="808080")
    for row in sheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_appendix_sheet(workbook: Workbook, rows: list[CombinedRecord]) -> None:
    sheet = workbook.active
    sheet.title = "Appendix Table"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:I1")
    sheet["A1"] = "Llama3-70B Throughput and Energy Efficiency"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:I2")
    sheet["A2"] = "Attention and FFN are combined; edge uses batch 1 and server uses batch 8."
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="555555")
    sheet["A2"].alignment = Alignment(horizontal="center")

    sheet.merge_cells("A4:A6")
    sheet["A4"] = "Sequence\nLength"
    sheet.merge_cells("B4:E4")
    sheet["B4"] = "Edge (Batch 1)"
    sheet.merge_cells("F4:I4")
    sheet["F4"] = "Server (Batch 8)"
    for start in (2, 6):
        sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=start + 1)
        sheet.cell(5, start, "Prefill")
        sheet.merge_cells(start_row=5, start_column=start + 2, end_row=5, end_column=start + 3)
        sheet.cell(5, start + 2, "Decode")
    for column in (2, 4, 6, 8):
        sheet.cell(6, column, "Throughput\n(GFLOPS)")
        sheet.cell(6, column + 1, "Energy Efficiency\n(GFLOPS/J)")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    subheader_fill = PatternFill("solid", fgColor="EAF2F8")
    for row in range(4, 7):
        for column in range(1, 10):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill if row == 4 else subheader_fill

    by_key = {(row.config, row.action, row.length): row for row in rows}
    lengths = sorted({row.length for row in rows})
    for offset, length in enumerate(lengths, 7):
        sheet.cell(offset, 1, length)
        keys = (
            ("edge", "prefill"), ("edge", "decode"),
            ("server", "prefill"), ("server", "decode"),
        )
        column = 2
        for config, action in keys:
            row = by_key[(config, action, length)]
            sheet.cell(offset, column, row.throughput_gflops)
            sheet.cell(offset, column + 1, row.efficiency_gflops_per_j)
            column += 2
        for cell in sheet[offset]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(2, 10):
            sheet.cell(offset, column).number_format = "0.00"

    sheet.column_dimensions["A"].width = 13
    for column in range(2, 10):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.row_dimensions[1].height = 23
    sheet.row_dimensions[2].height = 18
    for row in range(4, 7):
        sheet.row_dimensions[row].height = 28
    sheet.freeze_panes = "B7"
    apply_table_border(sheet, 4, 6 + len(lengths), 1, 9)
    sheet.print_area = f"A1:I{6 + len(lengths)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "Llama3-70B appendix data"


def write_detail_sheet(workbook: Workbook, rows: list[CombinedRecord]) -> None:
    sheet = workbook.create_sheet("Detailed Data")
    headers = [
        "Model", "Block", "Action", "Config", "Batch", "Length",
        "Attention latency (ms)", "FFN latency (ms)", "Total latency (ms)",
        "Attention energy (J)", "FFN energy (J)", "Total energy (J)",
        "Attention FLOPs", "FFN FLOPs", "Total FLOPs", "Throughput (GFLOPS)",
        "Energy efficiency (GFLOPS/J)",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row.model, row.block, row.action, row.config, row.batch, row.length,
                row.attention_latency_ms, row.ffn_latency_ms, row.latency_ms,
                row.attention_energy_j, row.ffn_energy_j, row.energy_j,
                row.attention_flops, row.ffn_flops, row.flops,
                row.throughput_gflops, row.efficiency_gflops_per_j,
            ]
        )
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
        for column in range(7, 13):
            row[column - 1].number_format = "0.000000"
        for column in range(13, 16):
            row[column - 1].number_format = "0"
        for column in range(16, 18):
            row[column - 1].number_format = "0.0000"
    widths = [16, 8, 11, 10, 8, 10, 21, 18, 18, 20, 16, 17, 18, 16, 18, 22, 30]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Q{sheet.max_row}"
    apply_table_border(sheet, 1, sheet.max_row, 1, len(headers))


def write_methodology_sheet(workbook: Workbook, source: Path) -> None:
    sheet = workbook.create_sheet("Methodology")
    sheet.sheet_view.showGridLines = False
    content = [
        ("Source", str(source)),
        ("Aggregation", "Attention and FFN metrics are added for each action, config, batch, and sequence length."),
        ("Total latency", "attention latency (ms) + FFN latency (ms)"),
        ("Total energy", "attention energy (J) + FFN energy (J)"),
        ("Total FLOPs", "attention FLOPs + FFN FLOPs"),
        ("Throughput", "total FLOPs / (total latency in ms * 1e6), reported as GFLOPS"),
        ("Energy efficiency", "total FLOPs / (total energy in J * 1e9), reported as GFLOPS/J"),
        ("Edge configuration", "Batch 1"),
        ("Server configuration", "Batch 8"),
    ]
    sheet.append(["Field", "Definition"])
    for item in content:
        sheet.append(item)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", size=10, bold=True)
        row[1].font = Font(name="Arial", size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 95
    apply_table_border(sheet, 1, sheet.max_row, 1, 2)


def write_workbook(rows: list[CombinedRecord], path: Path, source: Path) -> None:
    workbook = Workbook()
    write_appendix_sheet(workbook, rows)
    write_detail_sheet(workbook, rows)
    write_methodology_sheet(workbook, source)
    workbook.save(path)

    # Reopen once to catch malformed workbook output before reporting success.
    checked = load_workbook(path, data_only=False, read_only=True)
    if checked.sheetnames != ["Appendix Table", "Detailed Data", "Methodology"]:
        raise ValueError(f"unexpected workbook sheets: {checked.sheetnames}")
    checked.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--log", type=Path, default=Path("test/llama_70b.log"))
    parser.add_argument(
        "--output-dir", type=Path, default=Path("evaluation/results/llama3-70b")
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    rows = combine(parse_log(args.log))
    if len(rows) != 20:
        raise SystemExit(f"expected 20 combined workloads, found {len(rows)}")
    if {row.model for row in rows} != {"llama3-70b"}:
        raise SystemExit(f"unexpected models: {sorted({row.model for row in rows})}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(rows, args.output_dir / "llama3_70b_combined.csv")
    write_workbook(rows, args.output_dir / "llama3_70b_appendix.xlsx", args.log)
    print(f"Wrote {len(rows)} combined rows to {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
